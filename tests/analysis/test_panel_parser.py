"""Panel Parser 测试：多日期子表 -> 面板 JSON。"""

from __future__ import annotations

import sys
from pathlib import Path

from openpyxl import Workbook


PROJECT_ROOT = Path(__file__).resolve().parents[2]
SRC_PATH = PROJECT_ROOT / "src"
if str(SRC_PATH) not in sys.path:
    sys.path.insert(0, str(SRC_PATH))


def _write_stock_panel_xlsx(path: Path) -> None:
    wb = Workbook()
    ws1 = wb.active
    ws1.title = "2026-04-08"
    ws2 = wb.create_sheet("2026-04-09")
    wb.create_sheet("Notes")

    header = [
        "Symbol",
        "Description",
        "X_D_Trend_State",
        "X_D_State_Bars",
        "X_W_Trend_State",
        "X_W_State_Bars",
        "X_M_Trend_State",
        "X_M_State_Bars",
        "Price",
        "Volume",
    ]
    for ws in (ws1, ws2):
        ws.append(header)
        ws.append(["AAPL", "Apple", -1, 3, -1, 2, 1, 5, 100.0, 10_000])
        ws.append(["TSLA", "Tesla", 1, 2, 0, 1, 1, 7, 200.0, 20_000])

    wb.save(path)


def _write_etf_panel_xlsx(path: Path) -> None:
    wb = Workbook()
    ws1 = wb.active
    ws1.title = "2026-04-08"
    ws2 = wb.create_sheet("2026-04-09")
    wb.create_sheet("Notes")

    header = [
        "Symbol",
        "Description",
        "X_D_Trend_State",
        "X_D_State_Bars",
        "X_W_Trend_State",
        "X_W_State_Bars",
        "X_M_Trend_State",
        "X_M_State_Bars",
    ]
    for ws in (ws1, ws2):
        ws.append(header)
        ws.append([159326, "ETF-A", 1, 2, 1, 5, 1, 5])
        ws.append([159338, "ETF-B", 0, 6, 0, 1, 1, 5])

    wb.save(path)


def test_parse_panel_xlsx_for_stock_builds_series_and_date_range(tmp_path: Path) -> None:
    from analysis.panel_parser import parse_panel_xlsx  # noqa: E402

    file_path = tmp_path / "stock.xlsx"
    _write_stock_panel_xlsx(file_path)

    panel = parse_panel_xlsx(file_path=file_path, instrument_type="stock", max_sheets=0)

    assert panel["sheet_count"] == 2
    assert panel["date_range"] == {"start": "2026-04-08", "end": "2026-04-09"}
    assert panel["symbol_count"] == 2
    assert panel["data_completeness"]["missing_price_ratio"] == 0.0

    symbols = {item["symbol"]: item for item in panel["symbols"]}
    assert set(symbols.keys()) == {"AAPL", "TSLA"}
    aapl = symbols["AAPL"]
    assert aapl["description"] == "Apple"
    assert len(aapl["series"]) == 2
    assert aapl["series"][0]["date"] == "2026-04-08"
    assert aapl["series"][0]["price"] == 100.0


def test_parse_panel_xlsx_for_etf_sets_missing_price_ratio_and_keeps_series(tmp_path: Path) -> None:
    from analysis.panel_parser import parse_panel_xlsx  # noqa: E402

    file_path = tmp_path / "etf.xlsx"
    _write_etf_panel_xlsx(file_path)

    panel = parse_panel_xlsx(file_path=file_path, instrument_type="etf", max_sheets=0)

    assert panel["sheet_count"] == 2
    assert panel["date_range"] == {"start": "2026-04-08", "end": "2026-04-09"}
    assert panel["symbol_count"] == 2
    assert panel["data_completeness"]["missing_price_ratio"] == 1.0

    symbols = {item["symbol"]: item for item in panel["symbols"]}
    assert set(symbols.keys()) == {"159326", "159338"}
    etf_a = symbols["159326"]
    assert etf_a["description"] == "ETF-A"
    assert len(etf_a["series"]) == 2
    assert etf_a["series"][0]["date"] == "2026-04-08"
    assert etf_a["series"][0]["price"] is None

