"""Excel 解析与标准化模块。"""

from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from common.types import InstrumentRow, InstrumentType
from ingestion.input_contract_validator import ValidatedFilePair


class ExcelParseError(ValueError):
    """Excel 解析错误。"""


REQUIRED_COLUMNS = {
    "symbol": "Symbol",
    "description": "Description",
    "x_d_trend_state": "X_D_Trend_State",
    "x_d_state_bars": "X_D_State_Bars",
    "x_w_trend_state": "X_W_Trend_State",
    "x_w_state_bars": "X_W_State_Bars",
    "x_m_trend_state": "X_M_Trend_State",
    "x_m_state_bars": "X_M_State_Bars",
}


def _build_header_index(header_row: tuple[object, ...]) -> dict[str, int]:
    """构建首行标题到列索引的映射。"""
    mapping: dict[str, int] = {}
    for index, column_name in enumerate(header_row):
        if isinstance(column_name, str):
            mapping[column_name.strip()] = index
    return mapping


def _assert_required_columns(header_mapping: dict[str, int], file_path: Path) -> None:
    """校验 Excel 是否包含所有必填列。"""
    missing_columns = [name for name in REQUIRED_COLUMNS.values() if name not in header_mapping]
    if missing_columns:
        raise ExcelParseError(f"文件缺少必填列: {missing_columns}，文件: {file_path}")


def _parse_int(raw_value: Any, column_name: str, row_number: int, file_path: Path) -> int:
    """将单元格值解析为整数。"""
    if raw_value is None or raw_value == "":
        return 0
    try:
        return int(raw_value)
    except (TypeError, ValueError) as error:
        raise ExcelParseError(
            f"列 {column_name} 在第 {row_number} 行不是合法整数，文件: {file_path}"
        ) from error


def _parse_single_file(snapshot_date: str, file_path: Path, instrument_type: InstrumentType) -> list[InstrumentRow]:
    """解析单个 Excel 文件并输出统一 DTO 列表。"""
    workbook = load_workbook(file_path, read_only=True, data_only=True)
    try:
        worksheet = workbook.active
        header = next(worksheet.iter_rows(min_row=1, max_row=1, values_only=True), None)
        if header is None:
            return []
        header_mapping = _build_header_index(header)
        _assert_required_columns(header_mapping, file_path)

        parsed_rows: list[InstrumentRow] = []
        for row_number, row_values in enumerate(
            worksheet.iter_rows(min_row=2, values_only=True),
            start=2,
        ):
            symbol_value = row_values[header_mapping[REQUIRED_COLUMNS["symbol"]]]
            if symbol_value is None or str(symbol_value).strip() == "":
                continue

            parsed_rows.append(
                InstrumentRow(
                    snapshot_date=snapshot_date,
                    instrument_type=instrument_type,
                    symbol=str(symbol_value).strip(),
                    description=str(
                        row_values[header_mapping[REQUIRED_COLUMNS["description"]]] or ""
                    ).strip(),
                    x_d_trend_state=str(
                        row_values[header_mapping[REQUIRED_COLUMNS["x_d_trend_state"]]] or ""
                    ).strip(),
                    x_d_state_bars=_parse_int(
                        row_values[header_mapping[REQUIRED_COLUMNS["x_d_state_bars"]]],
                        REQUIRED_COLUMNS["x_d_state_bars"],
                        row_number,
                        file_path,
                    ),
                    x_w_trend_state=str(
                        row_values[header_mapping[REQUIRED_COLUMNS["x_w_trend_state"]]] or ""
                    ).strip(),
                    x_w_state_bars=_parse_int(
                        row_values[header_mapping[REQUIRED_COLUMNS["x_w_state_bars"]]],
                        REQUIRED_COLUMNS["x_w_state_bars"],
                        row_number,
                        file_path,
                    ),
                    x_m_trend_state=str(
                        row_values[header_mapping[REQUIRED_COLUMNS["x_m_trend_state"]]] or ""
                    ).strip(),
                    x_m_state_bars=_parse_int(
                        row_values[header_mapping[REQUIRED_COLUMNS["x_m_state_bars"]]],
                        REQUIRED_COLUMNS["x_m_state_bars"],
                        row_number,
                        file_path,
                    ),
                )
            )
        return parsed_rows
    finally:
        workbook.close()


def parse_validated_file_pair(validated_file_pair: ValidatedFilePair) -> tuple[list[InstrumentRow], list[InstrumentRow]]:
    """解析通过契约校验的文件对，返回 stock_rows 与 etf_rows。"""
    stock_rows = _parse_single_file(
        snapshot_date=validated_file_pair.snapshot_date,
        file_path=validated_file_pair.stock_path,
        instrument_type="stock",
    )
    etf_rows = _parse_single_file(
        snapshot_date=validated_file_pair.snapshot_date,
        file_path=validated_file_pair.etf_path,
        instrument_type="etf",
    )
    return stock_rows, etf_rows
