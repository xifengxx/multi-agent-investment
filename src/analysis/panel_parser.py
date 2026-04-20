"""Panel Analysis 面板解析模块：多日期子表 Excel -> 面板 JSON。"""

from __future__ import annotations

from dataclasses import dataclass
from datetime import date
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


@dataclass(frozen=True)
class _SheetRef:
    sheet_name: str
    sheet_date: date


def _try_parse_date(sheet_name: str) -> date | None:
    try:
        return date.fromisoformat(sheet_name.strip())
    except ValueError:
        return None


def _parse_int_maybe(value: Any) -> int | None:
    if value is None or value == "":
        return None
    if isinstance(value, bool):
        return None
    try:
        return int(value)
    except (TypeError, ValueError):
        return None


def _parse_float_maybe(value: Any) -> float | None:
    if value is None or value == "":
        return None
    if isinstance(value, bool):
        return None
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def _build_header_index(header_row: tuple[object, ...]) -> dict[str, int]:
    mapping: dict[str, int] = {}
    for idx, name in enumerate(header_row):
        if isinstance(name, str) and name.strip():
            mapping[name.strip()] = idx
    return mapping


def parse_panel_xlsx(*, file_path: Path, instrument_type: str, max_sheets: int) -> dict[str, Any]:
    """解析多日期子表 Excel 为面板 JSON。

    返回结构（核心字段）：
    - sheet_count, date_range, symbol_count
    - data_completeness（缺失率）
    - symbols: [{symbol, description, series:[{date, x_*, price, volume}, ...]}, ...]
    """
    workbook = load_workbook(file_path, read_only=True, data_only=True)
    try:
        sheet_refs: list[_SheetRef] = []
        for sheet_name in workbook.sheetnames:
            d = _try_parse_date(sheet_name)
            if d is None:
                continue
            sheet_refs.append(_SheetRef(sheet_name=sheet_name, sheet_date=d))

        sheet_refs.sort(key=lambda x: x.sheet_date)
        if max_sheets > 0 and len(sheet_refs) > max_sheets:
            sheet_refs = sheet_refs[-max_sheets:]

        if not sheet_refs:
            return {
                "sheet_count": 0,
                "date_range": {"start": "", "end": ""},
                "symbol_count": 0,
                "data_completeness": {
                    "missing_price_ratio": 1.0 if instrument_type == "etf" else 0.0,
                    "missing_volume_ratio": 1.0 if instrument_type == "etf" else 0.0,
                    "missing_trend_ratio": 1.0,
                },
                "symbols": [],
            }

        by_symbol: dict[str, dict[str, Any]] = {}

        total_points = 0
        missing_price = 0
        missing_volume = 0
        missing_trend = 0

        for ref in sheet_refs:
            ws = workbook[ref.sheet_name]
            header = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), None)
            if header is None:
                continue
            idx = _build_header_index(header)

            def col(name: str) -> int | None:
                return idx.get(name)

            i_symbol = col("Symbol")
            if i_symbol is None:
                continue

            i_desc = col("Description")
            i_xd = col("X_D_Trend_State")
            i_xd_bars = col("X_D_State_Bars")
            i_xw = col("X_W_Trend_State")
            i_xw_bars = col("X_W_State_Bars")
            i_xm = col("X_M_Trend_State")
            i_xm_bars = col("X_M_State_Bars")
            i_price = col("Price")
            i_volume = col("Volume")

            for row in ws.iter_rows(min_row=2, values_only=True):
                raw_symbol = row[i_symbol] if i_symbol < len(row) else None
                if raw_symbol is None or str(raw_symbol).strip() == "":
                    continue
                symbol = str(raw_symbol).strip()

                description = ""
                if i_desc is not None and i_desc < len(row):
                    description = str(row[i_desc] or "").strip()

                x_d = _parse_int_maybe(row[i_xd] if i_xd is not None and i_xd < len(row) else None)
                x_d_bars = _parse_int_maybe(
                    row[i_xd_bars] if i_xd_bars is not None and i_xd_bars < len(row) else None
                )
                x_w = _parse_int_maybe(row[i_xw] if i_xw is not None and i_xw < len(row) else None)
                x_w_bars = _parse_int_maybe(
                    row[i_xw_bars] if i_xw_bars is not None and i_xw_bars < len(row) else None
                )
                x_m = _parse_int_maybe(row[i_xm] if i_xm is not None and i_xm < len(row) else None)
                x_m_bars = _parse_int_maybe(
                    row[i_xm_bars] if i_xm_bars is not None and i_xm_bars < len(row) else None
                )

                price = _parse_float_maybe(row[i_price] if i_price is not None and i_price < len(row) else None)
                volume = _parse_float_maybe(
                    row[i_volume] if i_volume is not None and i_volume < len(row) else None
                )

                total_points += 1
                if price is None:
                    missing_price += 1
                if volume is None:
                    missing_volume += 1
                if x_d is None or x_w is None or x_m is None:
                    missing_trend += 1

                bucket = by_symbol.setdefault(
                    symbol,
                    {"symbol": symbol, "description": description, "series": []},
                )
                if description:
                    bucket["description"] = description
                bucket["series"].append(
                    {
                        "date": ref.sheet_date.isoformat(),
                        "x_d_trend_state": x_d,
                        "x_d_state_bars": int(x_d_bars or 0),
                        "x_w_trend_state": x_w,
                        "x_w_state_bars": int(x_w_bars or 0),
                        "x_m_trend_state": x_m,
                        "x_m_state_bars": int(x_m_bars or 0),
                        "price": price,
                        "volume": volume,
                    }
                )

        symbols: list[dict[str, Any]] = []
        for symbol in sorted(by_symbol.keys()):
            item = by_symbol[symbol]
            item["series"].sort(key=lambda x: str(x.get("date", "")))
            symbols.append(item)

        start_date = sheet_refs[0].sheet_date.isoformat()
        end_date = sheet_refs[-1].sheet_date.isoformat()

        denom = float(total_points) if total_points > 0 else 1.0
        return {
            "sheet_count": len(sheet_refs),
            "date_range": {"start": start_date, "end": end_date},
            "symbol_count": len(symbols),
            "data_completeness": {
                "missing_price_ratio": round(missing_price / denom, 6),
                "missing_volume_ratio": round(missing_volume / denom, 6),
                "missing_trend_ratio": round(missing_trend / denom, 6),
            },
            "symbols": symbols,
        }
    finally:
        workbook.close()

